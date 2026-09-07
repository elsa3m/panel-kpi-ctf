"""
Carga de los Excel del panel.

- MOV-FIBRA  -> tabla de KPIs por tienda y ejecutivo (hoja MOV-FIBRA), metas,
                fichas, EPA, BASE P EPA y CIERRE BONOS.
- FIBRA DRIVE-> solicitudes de fibra (AVANCE FIBRAS), resumen por tienda /
                ejecutivo (RESUMEN), evolutivo diario (EVOLUTIVO).
- ESCUCHAS   -> resultados de escuchas por ejecutivo (opcional).

Todas las funciones devuelven diccionarios / DataFrames "limpios" para que la
app no tenga que conocer la estructura del Excel.
"""
from __future__ import annotations

import datetime as dt
import re
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd

from . import columns as C
from .config import FUENTES, DATA_DIR


# ---------------------------------------------------------------------------
# utilidades
# ---------------------------------------------------------------------------
def _num(v, default=None):
    """Convierte a float; textos como '#REF!', '#N/A', '' -> default."""
    if v is None:
        return default
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return default
    s = str(v).strip().replace("%", "").replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return default


def _txt(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _fecha(v):
    if v is None or (isinstance(v, float) and pd.isna(v)) or v is pd.NaT:
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and 30000 < v < 80000:      # serial Excel
        return (dt.datetime(1899, 12, 30) + dt.timedelta(days=int(v))).date()
    if isinstance(v, str):
        for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return dt.datetime.strptime(v.strip()[:10], fmt).date()
            except ValueError:
                pass
    return None


def _pdv(v) -> str:
    """Código de tienda como texto sin decimales ('5003')."""
    n = _num(v)
    if n is not None and float(n).is_integer():
        return str(int(n))
    return _txt(v)


def _norm(s: str) -> str:
    """Normaliza encabezados para comparar (mayúsculas, sin acentos ni espacios extra)."""
    s = _txt(s).upper()
    s = re.sub(r"[ÁÀÄ]", "A", s)
    s = re.sub(r"[ÉÈË]", "E", s)
    s = re.sub(r"[ÍÌÏ]", "I", s)
    s = re.sub(r"[ÓÒÖ]", "O", s)
    s = re.sub(r"[ÚÙÜ]", "U", s)
    return re.sub(r"\s+", " ", s).strip()


def ruta_fuente(nombre: str) -> Path:
    return DATA_DIR / FUENTES[nombre]["archivo"]


def fuente_disponible(nombre: str) -> bool:
    return ruta_fuente(nombre).exists()


# ---------------------------------------------------------------------------
# MOV-FIBRA
# ---------------------------------------------------------------------------
@dataclass
class DatosMovFibra:
    fecha_corte: dt.date | None
    avance_esperado: float
    pesos: dict            # clave ficha -> peso
    topes: dict            # clave ficha -> tope
    estandares: dict       # clave KPI -> estándar mínimo
    tiendas: pd.DataFrame  # una fila por tienda
    ejecutivos: pd.DataFrame  # una fila por ejecutivo
    total: dict            # fila TOTAL CTF
    nombres: dict          # código ejecutivo -> nombre completo (hoja METAS)
    jornadas: dict         # código ejecutivo -> FT / PT (hoja CIERRE BONOS)
    epa_ejecutivo: pd.DataFrame
    encuestas: pd.DataFrame
    umbrales: dict = field(default_factory=dict)   # hoja CONV-CUMP: clave -> {"meta": x, "amarillo": y}
    archivo: str = ""
    avisos: list = field(default_factory=list)


def _leer_mov_fibra(ws, avisos: list) -> tuple[pd.DataFrame, dict]:
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    header = rows[C.HEADER_ROW - 1]

    # validar encabezados
    for clave, (col, esperado) in C.COLS.items():
        if esperado is None:
            continue
        real = header[col - 1] if col - 1 < len(header) else None
        if _norm(real) != _norm(esperado):
            avisos.append(f"MOV-FIBRA col {col} ({clave}): se esperaba '{esperado}' y se encontró '{_txt(real)}'.")

    registros = []
    for r in rows[C.FIRST_DATA_ROW - 1:]:
        pdv = r[C.COLS["pdv"][0] - 1]
        ejec = _txt(r[C.COLS["ejecutivo"][0] - 1])
        if pdv in (None, "") or ejec in ("", "0", "#N/A"):
            continue
        reg = {}
        for clave, (col, _) in C.COLS.items():
            v = r[col - 1] if col - 1 < len(r) else None
            if clave in ("pdv", "clave", "ejecutivo", "bf_estado"):
                reg[clave] = _txt(v)
            else:
                reg[clave] = _num(v)
        registros.append(reg)

    df = pd.DataFrame(registros)
    params = {
        "fecha_corte": _fecha(rows[C.CELL_FECHA_CORTE[0] - 1][C.CELL_FECHA_CORTE[1] - 1]),
        "avance_esperado": _num(rows[C.CELL_AVANCE_ESPERADO[0] - 1][C.CELL_AVANCE_ESPERADO[1] - 1], 0.0),
        "pesos": {k: _num(rows[C.ROW_PESOS_FICHA - 1][C.COLS[k][0] - 1], 0.0) for k, _ in C.FICHA_KPIS},
        "topes": {k: _num(rows[C.ROW_TOPES_FICHA - 1][C.COLS[k][0] - 1], 0.0) for k, _ in C.FICHA_KPIS},
        "estandares": {k: _num(rows[C.ROW_ESTANDARES - 1][col - 1]) for k, col in C.ESTANDARES.items()},
    }
    return df, params


def _leer_metas(wb) -> dict:
    """Código ejecutivo -> nombre completo (hoja METAS, col E) ."""
    nombres = {}
    if "METAS" not in wb.sheetnames:
        return nombres
    for r in wb["METAS"].iter_rows(min_row=5, values_only=True):
        cod = _txt(r[5]) if len(r) > 5 else ""
        nombre = _txt(r[4]) if len(r) > 4 else ""
        if cod and nombre and nombre != "#N/A" and cod != nombre:
            nombres[cod] = nombre
    return nombres


def _leer_jornadas(wb) -> dict:
    jornadas = {}
    if "CIERRE BONOS." not in wb.sheetnames:
        return jornadas
    for r in wb["CIERRE BONOS."].iter_rows(min_row=5, values_only=True):
        cod = _txt(r[1]) if len(r) > 1 else ""
        cargo = _txt(r[5]) if len(r) > 5 else ""
        if cod.startswith("CMA_") and cargo:
            jornadas[cod] = cargo
    return jornadas


def _leer_epa(wb) -> pd.DataFrame:
    """Tabla 'EPA <MES> POR EJECUTIVO' de la hoja EPA."""
    cols = ["pdv", "ejecutivo", "cep", "q_cep", "venta", "q_venta", "csim", "q_csim",
            "sstt", "q_sstt", "otros", "q_otros", "actitud", "q_total", "epa"]
    if "EPA" not in wb.sheetnames:
        return pd.DataFrame(columns=cols)
    rows = list(wb["EPA"].iter_rows(values_only=True))
    inicio = None
    for i, r in enumerate(rows):
        if len(r) > 4 and _norm(r[4]) == "USUARIO":
            inicio = i + 1
            break
    if inicio is None:
        return pd.DataFrame(columns=cols)
    out = []
    for r in rows[inicio:]:
        cod = _txt(r[4]) if len(r) > 4 else ""
        if not cod.startswith("CMA_"):
            continue
        out.append({
            "pdv": _txt(r[3]), "ejecutivo": cod,
            "cep": _num(r[6]), "q_cep": _num(r[7], 0), "venta": _num(r[8]), "q_venta": _num(r[9], 0),
            "csim": _num(r[10]), "q_csim": _num(r[11], 0), "sstt": _num(r[12]), "q_sstt": _num(r[13], 0),
            "otros": _num(r[14]), "q_otros": _num(r[15], 0), "actitud": _num(r[17]),
            "q_total": _num(r[18], 0), "epa": _num(r[21]) if len(r) > 21 else None,
        })
    return pd.DataFrame(out, columns=cols)


def _leer_encuestas(wb) -> pd.DataFrame:
    """Encuestas individuales (hoja BASE P EPA)."""
    cols = ["ejecutivo", "nota", "tienda", "fecha", "tipo_atencion", "literal"]
    if "BASE P EPA" not in wb.sheetnames:
        return pd.DataFrame(columns=cols)
    ws = wb["BASE P EPA"]
    rows = ws.iter_rows(values_only=True)
    header = [_norm(h) for h in next(rows)]

    def idx(*nombres):
        for n in nombres:
            if _norm(n) in header:
                return header.index(_norm(n))
        return None

    i_ej, i_nota = idx("Ejecutivo_hom", "Ejecutivo"), idx("Nota")
    i_tienda, i_fecha = idx("Nombre Tienda", "Tienda"), idx("Fecha Encuenta", "Fecha Encuesta")
    i_tipo, i_lit = idx("Tipo de Atención"), idx("Literal")
    out = []
    for r in rows:
        ej = _txt(r[i_ej]) if i_ej is not None and i_ej < len(r) else ""
        if not ej or ej == "0":
            continue
        out.append({
            "ejecutivo": ej,
            "nota": _num(r[i_nota]) if i_nota is not None else None,
            "tienda": _txt(r[i_tienda]) if i_tienda is not None else "",
            "fecha": _fecha(r[i_fecha]) if i_fecha is not None else None,
            "tipo_atencion": _txt(r[i_tipo]) if i_tipo is not None else "",
            "literal": _txt(r[i_lit]) if i_lit is not None else "",
        })
    return pd.DataFrame(out, columns=cols)


# columnas de la hoja CONV-CUMP -> clave interna del panel
CONV_CUMP_COLS = {
    6: "mov_conv", 7: "sus_conv", 8: "porta_conv", 9: "porta_peso", 10: "cvm_50", 11: "conv_fibra",
    12: "tasa_inst", 13: "pct_valid", 14: "pct_fact", 15: "eq_conv", 16: "att_eq_linea", 17: "att_seg",
    18: "acc_conv", 19: "ene_att", 20: "prot_att", 21: "epa",
    27: "mov_cump", 28: "sus_cump", 29: "l1_cump", 30: "l2_cump", 31: "porta_cump", 32: "fib_cump",
    33: "eq_cump", 34: "seg_cump", 35: "acc_cump",
}


def _leer_conv_cump(wb) -> dict:
    """Umbrales verde (fila 7 = META) y amarillo (fila 3) de la hoja CONV-CUMP."""
    if "CONV- CUMP" not in wb.sheetnames:
        return {}
    rows = list(wb["CONV- CUMP"].iter_rows(min_row=1, max_row=8, values_only=True))
    if len(rows) < 7:
        return {}
    out = {}
    for col, clave in CONV_CUMP_COLS.items():
        meta = _num(rows[6][col - 1]) if col - 1 < len(rows[6]) else None
        amar = _num(rows[2][col - 1]) if col - 1 < len(rows[2]) else None
        if meta is not None:
            out[clave] = {"meta": meta, "amarillo": amar if amar is not None else meta * 0.7}
    return out


def cargar_mov_fibra(path: Path | str) -> DatosMovFibra:
    path = Path(path)
    avisos: list[str] = []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if "MOV-FIBRA" not in wb.sheetnames:
        raise ValueError("El archivo no tiene la hoja 'MOV-FIBRA'.")

    df, params = _leer_mov_fibra(wb["MOV-FIBRA"], avisos)

    es_total = df["pdv"].str.upper().eq("CTF")
    es_tienda = (~es_total) & (df["ejecutivo"] == df["clave"])
    tiendas = df[es_tienda].copy()
    ejecutivos = df[(~es_total) & (~es_tienda)].copy()

    # nombre de tienda para cada ejecutivo
    nombre_tienda = dict(zip(tiendas["pdv"], tiendas["ejecutivo"]))
    ejecutivos["tienda"] = ejecutivos["pdv"].map(nombre_tienda).fillna("")
    tiendas["tienda"] = tiendas["ejecutivo"]

    # ejecutivos "activos": con meta móvil asignada
    ejecutivos["activo"] = ejecutivos["mov_meta"].fillna(0) > 0

    total = df[es_total].iloc[0].to_dict() if es_total.any() else {}

    datos = DatosMovFibra(
        fecha_corte=params["fecha_corte"],
        avance_esperado=params["avance_esperado"],
        pesos=params["pesos"],
        topes=params["topes"],
        estandares=params["estandares"],
        tiendas=tiendas.reset_index(drop=True),
        ejecutivos=ejecutivos.reset_index(drop=True),
        total=total,
        nombres=_leer_metas(wb),
        jornadas=_leer_jornadas(wb),
        epa_ejecutivo=_leer_epa(wb),
        encuestas=_leer_encuestas(wb),
        umbrales=_leer_conv_cump(wb),
        archivo=path.name,
        avisos=avisos,
    )
    wb.close()
    return datos


# ---------------------------------------------------------------------------
# FIBRA DRIVE
# ---------------------------------------------------------------------------
@dataclass
class DatosFibra:
    solicitudes: pd.DataFrame      # AVANCE FIBRAS (una fila por solicitud)
    resumen: pd.DataFrame          # RESUMEN (tiendas y ejecutivos)
    evolutivo: pd.DataFrame        # EVOLUTIVO (solicitudes por día)
    fecha_actualizacion: dt.date | None
    archivo: str = ""


def _df_desde_hoja(ws, header_row: int) -> pd.DataFrame:
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < header_row:
        return pd.DataFrame()
    header = [_norm(h) or f"COL{i+1}" for i, h in enumerate(rows[header_row - 1])]
    # encabezados duplicados -> sufijo
    vistos = {}
    for i, h in enumerate(header):
        if h in vistos:
            vistos[h] += 1
            header[i] = f"{h}_{vistos[h]}"
        else:
            vistos[h] = 1
    data = [list(r) + [None] * (len(header) - len(r)) for r in rows[header_row:]]
    df = pd.DataFrame(data, columns=header)
    return df.dropna(how="all")


def cargar_fibra(path: Path | str) -> DatosFibra:
    path = Path(path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    sol = _df_desde_hoja(wb["AVANCE FIBRAS"], 2) if "AVANCE FIBRAS" in wb.sheetnames else pd.DataFrame()
    if not sol.empty:
        sol = sol.rename(columns={
            "MES": "mes", "MES INST": "mes_inst", "FECHA DE SOLICITUD": "fecha_solicitud",
            "COD_SS": "pdv", "NOMBRE EJECUTIVO": "ejecutivo", "FECHA DE INSTALACION": "fecha_instalacion",
            "ESTADO": "estado", "INCLUYE TV": "incluye_tv", "ESTATUS": "estatus",
            "TIPO DE RECHAZO": "tipo_rechazo", "MOTIVO": "motivo",
        })
        sol["fecha_solicitud"] = sol["fecha_solicitud"].map(_fecha).astype(object)
        sol["fecha_instalacion"] = sol["fecha_instalacion"].map(_fecha).astype(object)
        sol["pdv"] = sol["pdv"].map(_pdv)
        sol["ejecutivo"] = sol["ejecutivo"].map(_txt)
        sol = sol[sol["ejecutivo"] != ""]
        # el RUT del cliente no se usa en el panel: se descarta al cargar
        sol = sol.drop(columns=[c for c in ("RUT",) if c in sol.columns])
        sol = sol.rename(columns={"ID": "id"})
        # intentos de agendamiento = fechas informadas (1RA / 2DA / 3DA)
        cols_fecha = [c for c in ("1RA FECHA", "2DA FECHA", "3DA FECHA") if c in sol.columns]
        if cols_fecha:
            sol["intentos"] = sol[cols_fecha].notna().sum(axis=1).clip(lower=1)
        else:
            sol["intentos"] = 1
        sol["reagendamientos"] = (sol["intentos"] - 1).clip(lower=0)
        sol["contratista"] = sol["CONTRATISTA PRODUCCION ENTEL"].map(_txt) if "CONTRATISTA PRODUCCION ENTEL" in sol.columns else ""
        sol["comuna"] = sol["COMUNA INGRESOS ENTEL"].map(_txt) if "COMUNA INGRESOS ENTEL" in sol.columns else ""
        sol["afinidad"] = sol["AFINIDAD"].map(_txt) if "AFINIDAD" in sol.columns else ""

    res = _df_desde_hoja(wb["RESUMEN"], 2) if "RESUMEN" in wb.sheetnames else pd.DataFrame()
    if not res.empty:
        res = res.rename(columns={
            "COD_SS": "pdv", "META SOLICIT.": "meta_sol", "DEBEN LLEVAR (SOLIC.)": "deben_sol",
            "SOLICITUDES OK": "sol_ok", "RECHAZO": "rechazo", "RECONT.": "recont",
            "META FIBRA": "meta_fibra", "REAL FIBRA NO AFINIDAD": "real_no_af",
            "REAL FIBRA AFINIDAD": "real_af", "REAL FIBRA": "real_fibra", "% CUMP.": "cump",
            "PROY.": "proy", "DEBEN LLEVAR INST.": "deben_inst", "META TV": "meta_tv",
            "REAL TV": "real_tv", "% CUMP._2": "cump_tv", "ATT": "att_tv",
        })
        res["pdv"] = res["pdv"].map(_pdv)
        # columna F (índice 5) trae el nombre de tienda o el código de ejecutivo
        res["nombre"] = res.iloc[:, 5].map(_txt)
        res["es_ejecutivo"] = res["nombre"].str.startswith("CMA_")
        for c in ("meta_sol", "deben_sol", "sol_ok", "rechazo", "recont", "meta_fibra",
                  "real_no_af", "real_af", "real_fibra", "cump", "deben_inst",
                  "meta_tv", "real_tv", "cump_tv", "att_tv"):
            if c in res.columns:
                res[c] = res[c].map(_num)
        res = res[res["nombre"] != ""]

    evo = _df_desde_hoja(wb["EVOLUTIVO"], 2) if "EVOLUTIVO" in wb.sheetnames else pd.DataFrame()

    fecha = None
    if not sol.empty and sol["fecha_solicitud"].notna().any():
        fechas = [f for f in sol["fecha_solicitud"] if isinstance(f, dt.date) and pd.notna(f)]
        fecha = max(fechas) if fechas else None
    wb.close()
    return DatosFibra(solicitudes=sol, resumen=res, evolutivo=evo, fecha_actualizacion=fecha, archivo=path.name)


# ---------------------------------------------------------------------------
# ESCUCHAS ENTEL  (export del Power BI "Adherencia KPIs Hogar por PDV - Ejecutivo")
# ---------------------------------------------------------------------------
ESCUCHAS_COLS = ["clave", "pdv", "nivel", "auditadas", "starlink", "latam_pass", "hogar",
                 "fibra_calidad", "fibra_estabilidad", "porta_motivo", "porta_objeciones", "porta_urgencia"]

# columna del export -> clave interna (se comparan normalizadas)
ESCUCHAS_MAPA = {
    "KPI FOCO N AUDITADAS": "auditadas",
    "STARLINK % TODAS": "starlink",
    "LATAM PASS % TODAS": "latam_pass",
    "MOTIVO HOGAR %": "hogar",
    "FIBRA CALIDAD %": "fibra_calidad",
    "FIBRA ESTABILIDAD %": "fibra_estabilidad",
    "MOTIVO PORTABILIDAD %": "porta_motivo",
    "PORTA OBJECIONES %": "porta_objeciones",
    "PORTA URGENCIA %": "porta_urgencia",
}


def _pdv_desde_texto(txt: str) -> str:
    """'5245 - Arauco Maipú' -> '5245'."""
    m = re.match(r"\s*(\d{3,6})\s*-", _txt(txt))
    return m.group(1) if m else ""


def cargar_escuchas(path: Path | str) -> pd.DataFrame:
    """
    Lee el Excel exportado desde Power BI (hoja 'Export' o la primera hoja).

    Cada fila queda clasificada en 'nivel':
      - "ejecutivo": agent_id = CMA_xxx
      - "tienda":    agent_id = Total  (el PDV trae el código de la tienda)
      - "ctf":       PDV = Total       (total de la franquicia)
      - "canal":     si el export incluye una fila de canal
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Export"] if "Export" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # fila de encabezados: la que contiene PDV / agent_id
    h_idx = None
    for i, r in enumerate(rows[:25]):
        norm = [_norm(v) for v in r]
        if "PDV" in norm and any(n in ("AGENT_ID", "AGENTE", "EJECUTIVO") for n in norm):
            h_idx = i
            break
    if h_idx is None:
        return pd.DataFrame(columns=ESCUCHAS_COLS)

    header = [_norm(v) for v in rows[h_idx]]
    i_pdv = header.index("PDV")
    i_ag = next((header.index(n) for n in ("AGENT_ID", "AGENTE", "EJECUTIVO") if n in header), None)
    pos = {clave: header.index(_norm(col)) for col, clave in ESCUCHAS_MAPA.items() if _norm(col) in header}

    out = []
    for r in rows[h_idx + 1:]:
        pdv_txt = _txt(r[i_pdv]) if i_pdv < len(r) else ""
        ag = _txt(r[i_ag]) if (i_ag is not None and i_ag < len(r)) else ""
        if not pdv_txt or pdv_txt.upper().startswith("FILTROS APLICADOS"):
            continue
        reg = {clave: _num(r[i]) for clave, i in pos.items() if i < len(r)}
        if ag.upper().startswith("CMA_"):
            nivel, clave = "ejecutivo", ag.upper()
        elif ag.upper() in ("TOTAL", "") and pdv_txt.upper() == "TOTAL":
            nivel, clave = "ctf", "CTF"
        elif ag.upper() == "CANAL" or pdv_txt.upper() == "CANAL":
            nivel, clave = "canal", "CANAL"
        elif ag.upper() == "TOTAL":
            nivel, clave = "tienda", _pdv_desde_texto(pdv_txt)
        else:
            continue
        reg.update({"clave": clave, "pdv": _pdv_desde_texto(pdv_txt), "nivel": nivel})
        out.append(reg)

    df = pd.DataFrame(out)
    for c in ESCUCHAS_COLS:
        if c not in df.columns:
            df[c] = None
    return df[ESCUCHAS_COLS]


# ---------------------------------------------------------------------------
# COLABORADORES  (hoja DOTACIÓN)
# ---------------------------------------------------------------------------
COLAB_COLS = ["ejecutivo", "nombre", "fecha_ingreso", "fecha_nacimiento", "jornada", "estatus"]

# encabezado esperado -> clave interna
COLAB_MAPA = {
    "IDENTIDAD RED": "ejecutivo",
    "NOMBRE COMPLETO": "nombre",
    "FECHA INGRESO": "fecha_ingreso",
    "FECHA NACIMIENTO": "fecha_nacimiento",
    "JORNADA": "jornada",
    "ESTADO ACTUAL": "estatus",
}


def cargar_colaboradores(path: Path | str) -> pd.DataFrame:
    """
    Lee la hoja DOTACIÓN y devuelve SOLO las columnas que el panel necesita.

    El archivo original trae datos personales que el panel no usa (RUT, dirección,
    teléfonos, contactos de emergencia, credenciales…). Esta función los descarta:
    nunca salen de aquí ni se guardan en disco.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["DOTACIÓN"] if "DOTACIÓN" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    # 1) formato reducido (el que guarda esta misma app): encabezados ya finales
    for i, r in enumerate(rows[:5]):
        header = [_norm(v) for v in r]
        if "EJECUTIVO" in header and "FECHA INGRESO".replace(" ", "_") in [x.replace(" ", "_") for x in header] or \
           ("EJECUTIVO" in header and "FECHA_INGRESO" in header):
            pos = {c: header.index(_norm(c)) for c in COLAB_COLS if _norm(c) in header}
            out = []
            for rr in rows[i + 1:]:
                cod = _txt(rr[pos["ejecutivo"]]).upper() if "ejecutivo" in pos and pos["ejecutivo"] < len(rr) else ""
                if not cod.startswith("CMA_"):
                    continue
                out.append({
                    "ejecutivo": cod,
                    "nombre": _txt(rr[pos["nombre"]]) if "nombre" in pos else "",
                    "fecha_ingreso": _fecha(rr[pos["fecha_ingreso"]]) if "fecha_ingreso" in pos else None,
                    "fecha_nacimiento": _fecha(rr[pos["fecha_nacimiento"]]) if "fecha_nacimiento" in pos else None,
                    "jornada": _txt(rr[pos["jornada"]]) if "jornada" in pos else "",
                    "estatus": _txt(rr[pos["estatus"]]) if "estatus" in pos else "",
                })
            return pd.DataFrame(out, columns=COLAB_COLS).drop_duplicates(subset=["ejecutivo"], keep="first")

    # 2) formato original de la planilla DOTACIÓN
    h_idx = None
    for i, r in enumerate(rows[:20]):
        if any(_norm(v) == "IDENTIDAD RED" for v in r):
            h_idx = i
            break
    if h_idx is None:
        return pd.DataFrame(columns=COLAB_COLS)

    header = [_norm(v) for v in rows[h_idx]]
    pos = {clave: header.index(_norm(col)) for col, clave in COLAB_MAPA.items() if _norm(col) in header}

    out = []
    for r in rows[h_idx + 1:]:
        i = pos.get("ejecutivo")
        cod = _txt(r[i]).upper() if (i is not None and i < len(r)) else ""
        if not cod.startswith("CMA_"):
            continue
        out.append({
            "ejecutivo": cod,
            "nombre": _txt(r[pos["nombre"]]) if "nombre" in pos else "",
            "fecha_ingreso": _fecha(r[pos["fecha_ingreso"]]) if "fecha_ingreso" in pos else None,
            "fecha_nacimiento": _fecha(r[pos["fecha_nacimiento"]]) if "fecha_nacimiento" in pos else None,
            "jornada": _txt(r[pos["jornada"]]) if "jornada" in pos else "",
            "estatus": _txt(r[pos["estatus"]]) if "estatus" in pos else "",
        })
    df = pd.DataFrame(out, columns=COLAB_COLS)
    return df.drop_duplicates(subset=["ejecutivo"], keep="first")


def guardar_colaboradores_reducido(origen_bytes: bytes, destino: Path) -> int:
    """
    Guarda en disco SOLO las columnas necesarias del archivo de dotación.

    Recibe el archivo subido en memoria, extrae las columnas del panel y escribe
    un Excel reducido. El archivo original nunca se guarda. Devuelve las filas.
    """
    import io
    df = cargar_colaboradores(io.BytesIO(origen_bytes))
    df.to_excel(destino, index=False, sheet_name="COLABORADORES")
    return len(df)
